"""Leitor da tabela de taxas contratadas (`02_tabela_taxas_contratadas.xlsx`).

Armadilhas de leitura encontradas na exploração:
- Duas abas: "Taxas vigentes" (usar) e "ANTIGO 2025 - nao usar" (nunca para taxa; é a
  única fonte de alguns dados cadastrais, ex.: CNPJ da Atlas Idiomas).
- "Taxas vigentes": R1–R2 título e "última atualização", R4 grupo mesclado (D4:I4),
  cabeçalho em R5, dados de R6 até a primeira linha vazia; depois de uma linha vazia,
  nota de rodapé ("* spread em branco = padrão").
- Percentual misturado na mesma coluna: número (0.0269) e texto ("2,69%", ex.: R20, R21).
- Ruído de float (0.007900000000000001): converter via str → Decimal e quantizar em
  4 casas de fração.
- CNPJ em 3 formatos: com máscara, só dígitos, e NÚMERO sem zero à esquerda (R14, R25).
- EC Orbe vazio em alguns lojistas (R8, R9 "EC pendente", R28) e nos só de remessa.
- Vigência em texto livre ("01/01/2026", "até 14/09/2026", "a partir de 15/09/2026").
- Spread remessa pode ser número, vazio ou o texto "padrão".
- Comentário de célula (A13) e linha inativa com estilo diferente (R30, obs "INATIVO").
- Ler com openpyxl (`data_only=True`); os comentários vêm do objeto da célula.

O cabeçalho é achado pelo conteúdo ("Lojista" e "CNPJ" na mesma linha) e as colunas pelo
nome, não por posição fixa. Linha que não dá para ler vira aviso, nunca chute.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from tarvo.formatos import FormatoInvalido, cnpj_valido, normalizar_cnpj, taxa_de_celula, texto_de_celula
from tarvo.leitores.orbe import ErroLeitura
from tarvo.modelo import Aviso, TaxaContratada

COLUNAS_CADASTRAIS = ("Lojista", "CNPJ", "EC Orbe", "Vigência", "Obs")


@dataclass
class LeituraTaxas:
    arquivo: str
    linhas: list[TaxaContratada] = field(default_factory=list)
    avisos: list[Aviso] = field(default_factory=list)


def ler_taxas_vigentes(caminho: Path, aba: str, colunas_taxa: list[str], casas: int) -> LeituraTaxas:
    """Lê a aba de taxas vigentes, com os valores crus e o número da linha da planilha.

    `aba`, `colunas_taxa` (cabeçalhos das colunas de taxa) e `casas` vêm de `regras.yaml`
    (taxas.aba_vigente, taxas.coluna_por_produto, taxas.casas_decimais_taxa).
    """
    nome = caminho.name
    wb = load_workbook(caminho, data_only=True)
    if aba not in wb.sheetnames:
        raise ErroLeitura(f"{nome}: aba {aba!r} não encontrada (abas: {wb.sheetnames})")
    ws = wb[aba]
    res = LeituraTaxas(arquivo=nome)

    linha_cab, cols = _achar_cabecalho(ws)
    if linha_cab is None:
        raise ErroLeitura(f"{nome}: cabeçalho com 'Lojista' e 'CNPJ' não encontrado na aba {aba!r}")
    faltando = [c for c in (*COLUNAS_CADASTRAIS, *colunas_taxa) if c not in cols]
    if faltando:
        raise ErroLeitura(f"{nome}: colunas não encontradas no cabeçalho R{linha_cab}: {faltando}")

    fim_do_bloco = False
    for r in range(linha_cab + 1, ws.max_row + 1):
        valores = {c: ws.cell(r, i).value for c, i in cols.items()}
        linha_vazia = all(texto_de_celula(ws.cell(r, i).value) is None for i in range(1, ws.max_column + 1))
        if linha_vazia:
            fim_do_bloco = True
            continue
        if fim_do_bloco:
            conteudo = " | ".join(t for i in range(1, ws.max_column + 1) if (t := texto_de_celula(ws.cell(r, i).value)))
            res.avisos.append(Aviso(nome, f"R{r}", f"linha depois do bloco de dados, ignorada (não é lojista): {conteudo!r}"))
            continue

        lojista = texto_de_celula(valores["Lojista"])
        if not lojista:
            res.avisos.append(Aviso(nome, f"R{r}", "linha sem nome de lojista, ignorada"))
            continue
        cnpj = normalizar_cnpj(valores["CNPJ"])
        if cnpj is None:
            res.avisos.append(Aviso(nome, f"R{r}", f"{lojista}: CNPJ ilegível {valores['CNPJ']!r}"))
        elif not cnpj_valido(cnpj):
            res.avisos.append(Aviso(nome, f"R{r}", f"{lojista}: CNPJ {cnpj} com dígito verificador inválido"))
            cnpj = None
        ec = texto_de_celula(valores["EC Orbe"])
        if ec is not None and not ec.isdigit():
            res.avisos.append(Aviso(nome, f"R{r}", f"{lojista}: EC Orbe não numérico {ec!r}, ignorado"))
            ec = None

        taxas: dict = {}
        for col in colunas_taxa:
            try:
                taxas[col] = taxa_de_celula(valores[col], casas)
            except FormatoInvalido as e:
                res.avisos.append(Aviso(nome, f"R{r}", f"{lojista}, coluna {col!r}: {e}"))
                taxas[col] = None

        comentario = ws.cell(r, cols["Lojista"]).comment
        res.linhas.append(
            TaxaContratada(
                linha_planilha=r,
                lojista=lojista,
                cnpj_bruto=repr(valores["CNPJ"]),
                cnpj=cnpj,
                ec_orbe=ec,
                taxas=taxas,
                vigencia_texto=texto_de_celula(valores["Vigência"]),
                obs=texto_de_celula(valores["Obs"]),
                comentario=comentario.text.strip() if comentario else None,
            )
        )
    return res


def _achar_cabecalho(ws) -> tuple[int | None, dict[str, int]]:
    for r in range(1, min(ws.max_row, 30) + 1):
        textos = {texto_de_celula(ws.cell(r, i).value): i for i in range(1, ws.max_column + 1)}
        if "Lojista" in textos and "CNPJ" in textos:
            return r, {t: i for t, i in textos.items() if t}
    return None, {}


def ler_aba_antiga(caminho: Path) -> list[dict]:
    """Lê a aba "ANTIGO 2025 - nao usar" só como referência cadastral (nunca como taxa)."""
    raise NotImplementedError
