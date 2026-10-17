import openpyxl,re
from decimal import Decimal as D, ROUND_HALF_UP, ROUND_HALF_UP
from collections import defaultdict
tx=openpyxl.load_workbook('/mnt/user-data/outputs/tarvo-fechamento/data/02_tabela_taxas_contratadas.xlsx')['Taxas vigentes']
sp={}; nm={}
for r in list(tx.iter_rows(values_only=True))[5:30]:
    if not r[0] or not r[1]: continue
    c=re.sub(r'\D','',str(r[1])).zfill(14); v=r[9]
    sp[c]=D('0.015') if v in (None,'padrão') else D(str(v)); nm[c]=r[0]
ws=openpyxl.load_workbook('ponte.xlsx')['Operações']
seen=set(); com=defaultdict(D); n=defaultdict(int); allsp=D(0); exc=[]
divA=defaultdict(D); divB=defaultdict(D); neg=[]
q=lambda x:x.quantize(D('0.01'),ROUND_HALF_UP)
for r in ws.iter_rows(min_row=6,max_row=238,values_only=True):
    op,dc,dl,cnpj,_,_,me,tc,ptax,vr,spr,spp=r[:12]
    c=re.sub(r'\D','',str(cnpj)).zfill(14); s=D(str(spr)); allsp+=s
    assert c in sp,(op,c)
    if op in seen: exc.append((op,'dup',s)); continue
    seen.add(op)
    assert dc.month==9
    me,tc,ptax=D(str(me)),D(str(tc)),D(str(ptax))
    if s<0: neg.append((op,nm[c],s, q(s-sp[c]*ptax*me))); continue
    com[c]+=q(s*D('0.5')); n[c]+=1
    # divergência
    a=(D(str(spp))-sp[c])*ptax*me
    b=s-sp[c]*ptax*me
    exact=(tc-ptax)/ptax-sp[c]
    if abs(exact)>D('0.00005'):
        divA[nm[c]]+=q(a); divB[nm[c]]+=q(b)
for c in sorted(com,key=lambda k:-com[k]): print(nm[c],com[c],n[c])
print('TOTAL',sum(com.values()),'ops',sum(n.values()),'lojistas',len(com))
print('Σspread',allsp,'exc',exc)
print('neg',neg,'Σneg spread',sum(x[2] for x in neg),'Σneg div',sum(x[3] for x in neg))
print('div via Spread%',dict(divA)); print('div via Spread R$',dict(divB))
print('50% neg',sum(x[2] for x in neg)/2,'50% dup',exc[0][2]/2,'50% do total Ponte',allsp/2)
