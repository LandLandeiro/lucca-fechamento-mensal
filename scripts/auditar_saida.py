"""Auditoria independente da saída (não importa o pacote `tarvo`).

Lê só os CSVs de `saida/` e as fontes brutas de `data/` (stdlib + openpyxl) e confere:
  F  formato contra docs/case.md, arquivo por arquivo (colunas na ordem, `;`, ponto decimal sem
     milhar, datas ISO, CNPJ 14 dígitos com DV, enums, confiança 0–1, UTF-8 sem BOM, LF);
  C  coerência entre arquivos (Σ notas = Σ comissões positivas, divergência aponta linha
     existente, pendência completa, todo id_saida no rastro com soma que fecha, toda linha do
     rastro existe na fonte).
Uso: python scripts/auditar_saida.py [--saida saida] [--dados data]. Sai com 1 se houver violação.
"""

from __future__ import annotations

import argparse
import csv
import email
import io
import re
import sys
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

import openpyxl

OBRIGATORIAS = {
    "comissoes_por_lojista.csv": ["id_saida", "cnpj", "lojista", "linha_negocio", "comissao_brl"],
    "divergencias_taxa.csv": ["id_saida", "cnpj", "linha_negocio", "ref_origem", "taxa_contratada_pct",
                              "taxa_aplicada_pct", "diferenca_brl"],
    "notas_a_emitir.csv": ["id_saida", "tomador_cnpj", "tomador", "linha_negocio", "valor_brl", "discriminacao"],
    "lancamentos_classificados.csv": ["id_saida", "banco", "ref_origem", "data", "descricao", "valor",
                                      "conta_codigo", "metodo", "confianca"],
    "pendencias.csv": ["id_saida", "tipo", "ref_origem", "valor_brl", "motivo", "quem_decide"],
    "rastro.csv": ["id_saida", "arquivo_origem", "ref_origem", "valor"],
}
NUMERO_REPORTADO = {
    "comissoes_por_lojista.csv": "comissao_brl", "divergencias_taxa.csv": "diferenca_brl",
    "notas_a_emitir.csv": "valor_brl", "lancamentos_classificados.csv": "valor", "pendencias.csv": "valor_brl",
}
ENUMS = {"linha_negocio": {"CHECKOUT", "REMESSA"}, "banco": {"AURORA", "MERIDIANO"}, "metodo": {"REGRA", "IA", "HUMANO"}}
RE_DINHEIRO = re.compile(r"^-?\d+\.\d{2}$")
RE_NUMERO = re.compile(r"^-?\d+(\.\d+)?$")
RE_VIRGULA_OU_MILHAR = re.compile(r"^-?\d{1,3}(\.\d{3})+(,\d+)?$|^-?\d+,\d+$")
RE_DATA_ISO = re.compile(r"^\d{4}-\d{2}-\d{2}$")
COLS_DINHEIRO = {"comissao_brl", "diferenca_brl", "valor_brl", "valor", "valor_total", "valor_contestado_brl",
                 "spread_brl", "diferenca_contrato_brl", "valor_bruto"}
COLS_CNPJ = {"cnpj", "tomador_cnpj", "lojista_cnpj"}


def cnpj_ok(c: str) -> bool:
    if not re.fullmatch(r"\d{14}", c or "") or len(set(c)) == 1:
        return False
    def dv(base, pesos):
        r = sum(int(d) * p for d, p in zip(base, pesos)) % 11
        return "0" if r < 2 else str(11 - r)
    p1 = [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]
    return c[12] == dv(c[:12], p1) and c[13] == dv(c[:13], [6] + p1)


F_CHECAGENS = {
    "F colunas": "colunas obrigatórias do case na ordem exata (6 arquivos)",
    "F separador": "';' como separador; toda linha com o mesmo nº de campos do cabeçalho",
    "F decimal": "dinheiro em 0.00 com ponto; _pct com ponto; nenhum valor com vírgula decimal ou milhar",
    "F datas": "colunas data/data_* em AAAA-MM-DD",
    "F CNPJ": "cnpj/tomador_cnpj/lojista_cnpj com 14 dígitos e DV válido; cnpj e tomador_cnpj nunca vazios",
    "F enums": "linha_negocio ∈ {CHECKOUT, REMESSA}; banco ∈ {AURORA, MERIDIANO}; metodo ∈ {REGRA, IA, HUMANO}",
    "F confiança": "confianca de 0 a 1 em REGRA/IA; vazia em HUMANO",
    "F encoding": "UTF-8 sem BOM",
    "F fim de linha": "só LF",
}


class Achados:
    def __init__(self):
        self.itens: dict[str, list[str]] = defaultdict(list)
        self.ok: dict[str, list[str]] = defaultdict(list)
        self.avisos: dict[str, list[str]] = defaultdict(list)

    def aviso(self, chave, msg):
        self.avisos[chave].append(msg)

    def erro(self, chave, msg):
        self.itens[chave].append(msg)

    def nota(self, chave, msg):
        self.ok[chave].append(msg)


# ---------------------------------------------------------------- fontes brutas

def fontes(dados: Path) -> dict[str, set[tuple[str, int]]]:
    """arquivo → {(ref_origem, linha)} lido direto dos arquivos do data pack."""
    out: dict[str, set[tuple[str, int]]] = {}
    orbe = next(dados.glob("01_orbe_*.csv"))
    linhas = orbe.read_text(encoding="utf-8").splitlines()
    out[orbe.name] = {(l.split(";")[0], n) for n, l in enumerate(linhas, start=1) if n > 1 and l.strip()}
    eml = next(dados.glob("03_email_*.eml"))
    msg = email.message_from_bytes(eml.read_bytes())
    anexo = next(p for p in msg.walk() if (p.get_filename() or "").endswith(".xlsx"))
    wb = openpyxl.load_workbook(io.BytesIO(anexo.get_payload(decode=True)), data_only=True)
    ws = next(w for w in wb.worksheets if any(str(c.value or "").startswith("Nº Opera") for c in w[5]))
    col = next(c.column for c in ws[5] if str(c.value or "").startswith("Nº Opera"))
    out[eml.name] = {(str(ws.cell(r, col).value), r) for r in range(6, ws.max_row + 1)
                     if re.fullmatch(r"PC-\d{4}-\d{4}", str(ws.cell(r, col).value or ""))}
    aurora = next(dados.glob("05a_*.csv"))
    out[aurora.name] = {(f"L{n}", n) for n, l in enumerate(aurora.read_text(encoding="utf-8").splitlines(), start=1)
                        if re.match(r"\d{2}/\d{2}/\d{4};", l) and "SALDO" not in l.split(";")[1]}
    ofx = next(dados.glob("05b_*.ofx"))
    linhas = ofx.read_text(encoding="latin-1").splitlines()
    refs, inicio = set(), None
    for n, l in enumerate(linhas, start=1):
        if l.strip().upper() == "<STMTTRN>":
            inicio = n
        m = re.match(r"\s*<FITID>(\S+)", l)
        if m and inicio:
            refs.add((m.group(1), inicio))
    out[ofx.name] = refs
    return out


def plano_lancavel(dados: Path) -> set[str]:
    p = next(dados.glob("04_plano_*.csv"))
    leitor = csv.DictReader(io.StringIO(p.read_text(encoding="latin-1")), delimiter=";")
    return {r["COD_CONTA"] for r in leitor if r["TIPO"] == "ANALITICA" and r["ACEITA_LANCTO"] == "S"}


# ---------------------------------------------------------------- F: formato

def ler(caminho: Path, a: Achados) -> list[dict] | None:
    bruto = caminho.read_bytes()
    nome = caminho.name
    if bruto.startswith(b"\xef\xbb\xbf"):
        a.erro("F encoding", f"{nome}: começa com BOM")
    if b"\r" in bruto:
        a.erro("F fim de linha", f"{nome}: tem CR (fim de linha não é só LF)")
    try:
        texto = bruto.decode("utf-8")
    except UnicodeDecodeError as e:
        a.erro("F encoding", f"{nome}: não é UTF-8 ({e})")
        return None
    linhas = list(csv.reader(io.StringIO(texto.removeprefix("\ufeff"), newline=""), delimiter=";"))
    if not linhas:
        a.erro("F separador", f"{nome}: vazio")
        return None
    cab = linhas[0]
    if len(cab) < 2:
        a.erro("F separador", f"{nome}: cabeçalho sem ';'")
    obrig = OBRIGATORIAS.get(nome)
    if obrig and cab[:len(obrig)] != obrig:
        a.erro("F colunas", f"{nome}: colunas obrigatórias fora da ordem: {cab[:len(obrig)]} ≠ {obrig}")
    tortas = [i for i, l in enumerate(linhas[1:], start=2) if len(l) != len(cab)]
    if tortas:
        a.erro("F separador", f"{nome}: {len(tortas)} linhas com nº de campos ≠ cabeçalho (ex.: {tortas[:3]})")
    return [dict(zip(cab, l)) for l in linhas[1:]]


def formato(nome: str, linhas: list[dict], a: Achados) -> None:
    for i, l in enumerate(linhas, start=2):
        local = f"{nome}:{i}"
        for col, v in l.items():
            if v and RE_VIRGULA_OU_MILHAR.match(v) and col not in ("descricao", "motivo", "discriminacao", "nota"):
                a.erro("F decimal", f"{local} {col}={v!r}: vírgula decimal ou separador de milhar")
            if col in COLS_DINHEIRO and v != "" and not RE_DINHEIRO.match(v):
                a.erro("F decimal", f"{local} {col}={v!r} fora de 0.00")
            if col.endswith("_pct") and v != "" and not RE_NUMERO.match(v):
                a.erro("F decimal", f"{local} {col}={v!r} não é número com ponto")
            if (col == "data" or col.startswith("data_")) and v != "" and not RE_DATA_ISO.match(v):
                a.erro("F datas", f"{local} {col}={v!r} fora de AAAA-MM-DD")
            if col in COLS_CNPJ and v != "" and not cnpj_ok(v):
                a.erro("F CNPJ", f"{local} {col}={v!r} não é CNPJ de 14 dígitos com DV válido")
            if col in ENUMS and v not in ENUMS[col]:
                a.erro("F enums", f"{local} {col}={v!r} fora de {sorted(ENUMS[col])}")
        if nome == "lancamentos_classificados.csv":
            c, m = l.get("confianca", ""), l.get("metodo")
            if m == "HUMANO":
                if c != "":
                    a.erro("F confiança", f"{local}: HUMANO com confiança {c!r}")
            elif not (RE_NUMERO.match(c or "") and Decimal("0") <= Decimal(c) <= Decimal("1")):
                a.erro("F confiança", f"{local}: {m} com confiança {c!r} fora de 0–1")
        for col in ("cnpj", "tomador_cnpj"):
            if col in l and l[col] == "":
                a.erro("F CNPJ", f"{local}: {col} vazio")
    # aviso (não é violação: o campo vai entre aspas, RFC 4180): texto com ';' quebra quem divide a linha por ';'
    obrig = set(OBRIGATORIAS.get(nome, []))
    cont = defaultdict(int)
    for l in linhas:
        for col, v in l.items():
            if ";" in v or "\n" in v:
                cont[col] += 1
    for col, n in sorted(cont.items()):
        onde = "coluna OBRIGATÓRIA (split(';') desloca as seguintes)" if col in obrig else "coluna extra, depois das obrigatórias"
        a.aviso("F aspas", f"{nome}: {n} valores de {col!r} com ';' entre aspas ({onde})")


# ---------------------------------------------------------------- C: coerência

def _num(v) -> Decimal:
    """Número da saída para somar; fora do formato conta 0 (a violação de formato já é reportada em F)."""
    if isinstance(v, (int, Decimal)):
        return Decimal(v)
    return Decimal(v) if RE_NUMERO.match(v or "") else Decimal(0)


def coerencia(d: dict[str, list[dict]], src: dict[str, set[tuple[str, int]]], lancaveis: set[str], a: Achados) -> None:
    D = _num
    com, notas, div, pen, lan, ras = (d[n] for n in ("comissoes_por_lojista.csv", "notas_a_emitir.csv",
                                                       "divergencias_taxa.csv", "pendencias.csv",
                                                       "lancamentos_classificados.csv", "rastro.csv"))
    # C1 notas × comissões
    for ln in ("CHECKOUT", "REMESSA"):
        pos = {c["cnpj"]: D(c["comissao_brl"]) for c in com if c["linha_negocio"] == ln and D(c["comissao_brl"]) > 0}
        nf = [n for n in notas if n["linha_negocio"] == ln]
        snf, scom = sum((D(n["valor_brl"]) for n in nf), D(0)), sum(pos.values(), D(0))
        if snf != scom:
            a.erro("C1 notas = comissões", f"{ln}: Σ notas {snf} ≠ Σ comissões positivas {scom}")
        else:
            a.nota("C1 notas = comissões", f"{ln}: Σ notas {snf} = Σ comissões positivas {scom} ({len(nf)} notas, {len(pos)} comissões > 0)")
        if ln == "CHECKOUT":
            for n in nf:
                if pos.get(n["tomador_cnpj"]) != D(n["valor_brl"]):
                    a.erro("C1 notas = comissões", f"{n['id_saida']}: {n['valor_brl']} ≠ comissão {pos.get(n['tomador_cnpj'])}")
            sem_nota = set(pos) - {n["tomador_cnpj"] for n in nf}
            if sem_nota:
                a.erro("C1 notas = comissões", f"comissão positiva sem nota: {sorted(sem_nota)}")
    neg = [c for c in com if D(c["comissao_brl"]) <= 0]
    a.nota("C1 notas = comissões", f"comissões ≤ 0 sem nota: {[(c['lojista'], c['comissao_brl']) for c in neg]}")

    # C2 divergência → linha existente
    refs_por_arquivo = {arq: {r for r, _ in s} for arq, s in src.items()}
    orbe = next(k for k in src if k.startswith("01_"))
    ponte = next(k for k in src if k.startswith("03_"))
    for x in div:
        arq = orbe if x["linha_negocio"] == "CHECKOUT" else ponte
        if x["ref_origem"] not in refs_por_arquivo[arq]:
            a.erro("C2 divergência aponta linha existente", f"{x['id_saida']}: {x['ref_origem']} não existe em {arq}")
    a.nota("C2 divergência aponta linha existente", f"{len(div)} divergências; ref_origem conferida contra {orbe} (CHECKOUT) e o anexo de {ponte} (REMESSA)")

    # C3 pendência completa
    for p in pen:
        vazios = [c for c in ("tipo", "ref_origem", "motivo", "quem_decide") if not (p.get(c) or "").strip()]
        if vazios:
            a.erro("C3 pendência completa", f"{p['id_saida']}: vazio em {vazios}")
    a.nota("C3 pendência completa", f"{len(pen)} pendências; {sum(1 for p in pen if not p['valor_brl'])} sem valor_brl")

    # C4 todo id_saida no rastro, soma fecha
    soma = defaultdict(lambda: D(0))
    for r in ras:
        soma[r["id_saida"]] += D(r["valor"])
    reportados = {}
    for nome, col in NUMERO_REPORTADO.items():
        for x in d[nome]:
            if x["id_saida"] in reportados:
                a.erro("C4 rastro fecha", f"{x['id_saida']} repetido entre arquivos")
            reportados[x["id_saida"]] = x[col]
    sem_valor = [i for i, v in reportados.items() if v == ""]
    for i, v in reportados.items():
        if v == "":
            continue
        if i not in soma:
            a.erro("C4 rastro fecha", f"{i}: {v} sem linha no rastro")
        elif soma[i] != D(v):
            a.erro("C4 rastro fecha", f"{i}: reportado {v} ≠ Σ rastro {soma[i]}")
    orfaos = sorted(i for i in soma if i not in reportados and not i.startswith("EXC-"))
    if orfaos:
        a.erro("C4 rastro fecha", f"ids no rastro sem número de saída: {orfaos[:5]}")
    a.nota("C4 rastro fecha", f"{len(reportados) - len(sem_valor)} números conferidos ({len(sem_valor)} sem valor); "
                              f"{len(ras)} linhas de rastro; {sum(1 for i in soma if i.startswith('EXC-'))} ids EXC- (exclusão anotada, sem número)")
    if "contestacao_ponte.csv" in d:
        cs = defaultdict(lambda: D(0))
        for x in d["contestacao_ponte.csv"]:
            cs[x["id_pendencia"]] += D(x["valor_contestado_brl"])
        for pid, v in cs.items():
            if reportados.get(pid) in (None, "") or D(reportados[pid]) != v:
                a.erro("C4 rastro fecha", f"contestacao_ponte.csv {pid}: Σ {v} ≠ pendência {reportados.get(pid)}")
        a.nota("C4 rastro fecha", f"contestacao_ponte.csv: {len(cs)} pendências, Σ por pendência = valor da pendência")
    if "sugestoes_regra.csv" in d:
        val = {(l["banco"], l["ref_origem"]): D(l["valor"]) for l in lan}
        for s in d["sugestoes_regra.csv"]:
            t = sum((val[(s["banco"], r.strip())] for r in s["refs"].split(",")), D(0))
            if t != D(s["valor_total"]):
                a.erro("C4 rastro fecha", f"sugestoes_regra.csv {s['id_sugestao']}: valor_total {s['valor_total']} ≠ Σ lançamentos {t}")
        a.nota("C4 rastro fecha", f"sugestoes_regra.csv: valor_total = Σ dos lançamentos citados em refs")

    # C5 toda linha do rastro existe na fonte
    pseudo, inexistentes = [], []
    for r in ras:
        s = src.get(r["arquivo_origem"])
        if s is None:
            inexistentes.append(f"{r['id_saida']}: arquivo {r['arquivo_origem']!r} não é do data pack")
            continue
        n = int(r["linha_arquivo"]) if (r.get("linha_arquivo") or "").isdigit() else 0
        if (r["ref_origem"], n) not in s:
            (pseudo if n == 0 else inexistentes).append(f"{r['id_saida']} {r['ref_origem']} L{n}")
    for x in inexistentes:
        a.erro("C5 rastro aponta linha de origem", x)
    for x in pseudo:
        a.erro("C5 rastro aponta linha de origem", f"linha sem origem física (linha 0): {x}")
    a.nota("C5 rastro aponta linha de origem", f"{len(ras) - len(pseudo) - len(inexistentes)} de {len(ras)} linhas do rastro existem na fonte (ref + linha)")

    # extra: conta lançável
    for l in lan:
        if l["conta_codigo"] and l["conta_codigo"] not in lancaveis:
            a.erro("C6 conta lançável", f"{l['id_saida']}: conta {l['conta_codigo']} não é ANALITICA com ACEITA_LANCTO=S")
    a.nota("C6 conta lançável", f"{sum(1 for l in lan if l['conta_codigo'])} lançamentos com conta; plano lido direto do 04_plano")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--saida", type=Path, default=Path("saida"))
    ap.add_argument("--dados", type=Path, default=Path("data"))
    args = ap.parse_args()
    a = Achados()
    dados = {}
    for caminho in sorted(args.saida.glob("*.csv")):
        linhas = ler(caminho, a)
        if linhas is not None:
            dados[caminho.name] = linhas
            formato(caminho.name, linhas, a)
            a.nota("F arquivos", f"{caminho.name}: {len(linhas)} linhas")
    faltando = set(OBRIGATORIAS) - set(dados)
    if faltando:
        a.erro("F arquivos", f"faltam {sorted(faltando)}")
    else:
        coerencia(dados, fontes(args.dados), plano_lancavel(args.dados), a)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    for chave in F_CHECAGENS:
        if chave not in a.itens:
            a.nota(chave, F_CHECAGENS[chave])
    for chave in sorted(set(a.itens) | set(a.ok) | set(a.avisos)):
        status = "❌" if a.itens.get(chave) else ("⚠️" if a.avisos.get(chave) else "✅")
        print(f"{status} {chave}")
        for m in a.avisos.get(chave, []):
            print(f"    AVISO {m}")
        for m in a.ok.get(chave, []):
            print(f"    {m}")
        for m in a.itens.get(chave, [])[:20]:
            print(f"    ERRO {m}")
        if len(a.itens.get(chave, [])) > 20:
            print(f"    ... e mais {len(a.itens[chave]) - 20}")
    total = sum(len(v) for v in a.itens.values())
    print(f"\nviolações: {total}")
    return 1 if total else 0


if __name__ == "__main__":
    sys.exit(main())
